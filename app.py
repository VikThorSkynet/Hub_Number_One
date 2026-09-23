import csv
import io
import json
import os
import re
import sqlite3
import unicodedata
from pathlib import Path
from datetime import datetime, date
from urllib.parse import urlsplit
from flask import Flask, request, jsonify, send_from_directory, abort, send_file
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = Path(os.environ.get('AGENDA_DATA', str(Path(os.environ.get('LOCALAPPDATA', str(Path.home()))) / 'AgendaTelefonica')))
if str(DATA).startswith('\\\\'):
    raise RuntimeError('AGENDA_DATA deve ser uma pasta local do servidor, não uma pasta de rede.')
DATA.mkdir(parents=True, exist_ok=True)
app = Flask(__name__, static_folder='static')
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024

def norm(s):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(s or '').lower()) if not unicodedata.combining(c))

def phone(value):
    raw = str(value).strip()
    digits = re.sub(r'\D', '', raw)
    national = digits
    if national.startswith('0055'):
        national = national[4:]
    elif national.startswith('55') and (raw.startswith('+') or len(national) >= 12):
        national = national[2:]
    if national.startswith('0') and len(national) in (11, 12):
        national = national[1:]
    local = national[2:] if len(national) in (10, 11) else national
    reason = 'Menos de 9 dígitos sem DDD e país' if len(local) < 9 else ''
    if len(national) > 11 or (raw.startswith('+') and not digits.startswith('55')):
        reason = 'Formato internacional ou comprimento não reconhecido; revisar'
    return dict(value=raw, digits=digits, national=national, local=local, issue=reason)

class Connection(sqlite3.Connection):
    def __exit__(self, *args):
        try:
            return super().__exit__(*args)
        finally:
            self.close()

def db():
    con = sqlite3.connect(DATA / 'agenda.sqlite3', timeout=30, factory=Connection)
    con.row_factory = sqlite3.Row
    con.execute('PRAGMA foreign_keys=ON')
    return con

def init():
    with db() as c:
        c.execute('PRAGMA journal_mode=WAL')
        c.executescript('''
        CREATE TABLE IF NOT EXISTS contacts(id INTEGER PRIMARY KEY, data TEXT NOT NULL, search TEXT NOT NULL, revision INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE IF NOT EXISTS records(source TEXT NOT NULL, identity TEXT NOT NULL, contact_id INTEGER NOT NULL REFERENCES contacts(id), PRIMARY KEY(source,identity));
        CREATE TABLE IF NOT EXISTS imports(id INTEGER PRIMARY KEY, source TEXT, filename TEXT, created TEXT, added INTEGER, updated INTEGER, skipped INTEGER);
        CREATE TABLE IF NOT EXISTS hub_links(id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, url TEXT NOT NULL, description TEXT NOT NULL DEFAULT '');
        CREATE TABLE IF NOT EXISTS hub_notices(id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, body TEXT NOT NULL, color TEXT NOT NULL, seconds INTEGER NOT NULL, revision INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE IF NOT EXISTS hub_events(id INTEGER PRIMARY KEY AUTOINCREMENT, event_date TEXT NOT NULL, title TEXT NOT NULL, category TEXT NOT NULL, description TEXT NOT NULL DEFAULT '', source TEXT NOT NULL DEFAULT 'manual', source_key TEXT UNIQUE, revision INTEGER NOT NULL DEFAULT 1);
        CREATE INDEX IF NOT EXISTS idx_hub_events_date ON hub_events(event_date);
        CREATE TABLE IF NOT EXISTS hub_classes(id INTEGER PRIMARY KEY AUTOINCREMENT, weekday INTEGER NOT NULL, room INTEGER NOT NULL, start_time TEXT NOT NULL, course TEXT NOT NULL, students INTEGER, teacher TEXT NOT NULL DEFAULT '', revision INTEGER NOT NULL DEFAULT 1, UNIQUE(weekday,room,start_time));
        ''')

def prepare(d):
    out = {k: str(d.get(k, '') or '').strip() for k in ('name','company','title','department','notes','emails')}
    values = d.get('phones', [])
    if isinstance(values, str): values = re.split(r'[;\n]', values)
    out['phones'] = list(dict.fromkeys(str(x).strip() for x in values if str(x).strip()))
    out['sources'] = d.get('sources', [])
    out['manual'] = bool(d.get('manual', False))
    if not out['name']: out['name'] = next(iter(out['phones']), '') or out['emails']
    return out

def save(c, d, cid=None):
    search = norm(' '.join(str(v) for v in d.values())) + ' ' + ' '.join(phone(p)['digits']+' '+phone(p)['national'] for p in d['phones'])
    if cid:
        c.execute('UPDATE contacts SET data=?,search=?,revision=revision+1 WHERE id=?', (json.dumps(d,ensure_ascii=False),search,cid))
    else:
        cid = c.execute('INSERT INTO contacts(data,search) VALUES(?,?)',(json.dumps(d,ensure_ascii=False),search)).lastrowid
    return cid

def parse_file(content, filename):
    if filename.lower().endswith('.xlsx'):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = list(wb.active.values)
        wb.close()
        if not rows: raise ValueError('Planilha vazia.')
        headers = [str(v or '').strip() for v in rows[0]]
        raw = [dict(zip(headers, [str(v) if v is not None else '' for v in row])) for row in rows[1:]]
    elif filename.lower().endswith('.csv'):
        try: text = content.decode('utf-8-sig')
        except UnicodeDecodeError: text = content.decode('cp1252')
        try: dialect = csv.Sniffer().sniff(text[:10000], delimiters=',;\t')
        except csv.Error: dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        headers = reader.fieldnames or []
        raw = list(reader)
    else: raise ValueError('Use um arquivo CSV ou XLSX.')
    normalized = [norm(h) for h in headers]
    if not any(h in ('name','nome','first name','given name','telefone','celular') or re.match(r'phone \d+ - value',h) for h in normalized):
        raise ValueError('Cabeçalhos não reconhecidos. Use exportação Google CSV ou colunas Nome, Telefone, E-mail, Empresa, Cargo, Setor e Observações.')
    result = []
    for row in raw:
        r = {norm(k).strip(): str(v or '').strip() for k,v in row.items() if k}
        def get(*keys): return next((r[k] for k in keys if r.get(k)), '')
        name = get('name','nome') or ' '.join(filter(None,[get('first name','given name'),get('middle name','additional name'),get('last name','family name')]))
        phones = []
        for k,v in r.items():
            if re.match(r'phone \d+ - value$', k) or k in ('telefone','telefones','celular','ramal'):
                phones.extend(re.split(r'\s*:::?\s*|[;\n]', v))
        emails = '; '.join(v for k,v in r.items() if re.match(r'e-mail \d+ - value$',k) or k in ('email','e-mail'))
        d = prepare(dict(name=name,phones=phones,emails=emails,company=get('organization name','organization 1 - name','empresa'),title=get('organization title','organization 1 - title','cargo'),department=get('organization department','organization 1 - department','setor','departamento'),notes=get('notes','observacoes','observacao')))
        if d['name']: result.append(d)
    if not result: raise ValueError('Nenhum contato encontrado.')
    return result

def identity(d):
    return norm(d['name']) + '|' + '|'.join(sorted(phone(p)['national'] for p in d['phones'])) + '|' + norm(d['emails'])

def import_contacts(rows, source, filename):
    added = updated = skipped = 0
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        existing = [(r['id'],json.loads(r['data'])) for r in c.execute('SELECT * FROM contacts')]
        by_id = dict(existing)
        index = {}
        for cid,d in existing:
            index.setdefault(norm(d['name']), []).append(cid)
        for d in rows:
            key = identity(d)
            match = c.execute('SELECT contact_id FROM records WHERE source=? AND identity=?',(source,key)).fetchone()
            cid = match[0] if match else None
            if not cid:
                candidates=[]
                for candidate in index.get(norm(d['name']), []):
                    old=by_id[candidate]
                    shared = {phone(p)['national'] for p in old['phones']} & {phone(p)['national'] for p in d['phones']}
                    email_shared = set(norm(old['emails']).split('; ')) & set(norm(d['emails']).split('; ')) - {''}
                    if shared or email_shared: candidates.append(candidate)
                if len(candidates)==1: cid=candidates[0]
            if cid:
                old=by_id[cid]
                merged=dict(old)
                merged['sources']=list(dict.fromkeys(old['sources']+[source]))
                if not old['manual']:
                    for k in ('name','company','title','department','notes'):
                        if d[k]: merged[k]=d[k]
                    merged['phones']=list(dict.fromkeys(old['phones']+d['phones']))
                    merged['emails']='; '.join(dict.fromkeys(filter(None,old['emails'].split('; ')+d['emails'].split('; '))))
                if merged != old: save(c,merged,cid); updated+=1
                else: skipped+=1
                by_id[cid]=merged
            else:
                d['sources']=[source]
                cid=save(c,d)
                by_id[cid]=d
                index.setdefault(norm(d['name']),[]).append(cid)
                added+=1
            c.execute('INSERT OR IGNORE INTO records VALUES(?,?,?)',(source,key,cid))
        c.execute('INSERT INTO imports(source,filename,created,added,updated,skipped) VALUES(?,?,?,?,?,?)',(source,filename,datetime.now().isoformat(timespec='seconds'),added,updated,skipped))
    return dict(added=added,updated=updated,skipped=skipped,total=len(rows))

@app.before_request
def same_origin():
    if request.method in ('POST','PUT','DELETE'):
        origin=request.headers.get('Origin')
        if origin and origin != request.host_url.rstrip('/'): abort(403)

@app.get('/')
def home(): return send_from_directory(app.static_folder,'hub.html')

@app.get('/lista-telefonica')
def directory(): return send_from_directory(app.static_folder,'index.html')

@app.get('/calculadora')
def calculator(): return send_from_directory(app.static_folder,'calculadora.html')

@app.get('/links')
def links_page(): return send_from_directory(app.static_folder,'links.html')

@app.get('/calendario')
def calendar_page(): return send_from_directory(app.static_folder,'calendario.html')

@app.get('/mapa-de-turmas')
def classes_page(): return send_from_directory(app.static_folder,'turmas.html')

EVENT_CATEGORIES = {'prova','notas','feriado','aulas','segunda_chamada','projeto','sexta_letiva','encerramento','outro'}
SHEET_COLORS = {
    'FFCCFFFF': ('prova','Semana de provas'),
    'FFFFFF9B': ('notas','Envio de notas e boletins'),
    'FFCC66FF': ('sexta_letiva','Sexta letiva'),
    'FFFF99FF': ('aulas','Início das aulas'),
    'FF7F7F7F': ('feriado','Feriado ou recesso'),
    'FFBFBFBF': ('feriado','Feriado ou recesso'),
    'FF92D050': ('segunda_chamada','2ª chamada de provas'),
    'FFFBCDA7': ('projeto','Book Project Week'),
}

def valid_date(value):
    if not isinstance(value,str) or not re.fullmatch(r'\d{4}-\d{2}-\d{2}',value): return False
    try: date.fromisoformat(value)
    except ValueError: return False
    return True

def event_payload(data):
    if not isinstance(data,dict): return None
    event_date,title,category,description=[data.get(k,'') for k in ('event_date','title','category','description')]
    if not valid_date(event_date) or not isinstance(title,str) or not 1<=len(title.strip())<=120 or not isinstance(category,str) or category not in EVENT_CATEGORIES or not isinstance(description,str) or len(description)>500: return None
    return event_date,title.strip(),category,description.strip()

@app.get('/api/events')
def list_events():
    start,end=request.args.get('from'),request.args.get('to')
    if (start and not valid_date(start)) or (end and not valid_date(end)) or (start and end and start>end): return jsonify(error='Período inválido.'),400
    clauses=[]; params=[]
    if start: clauses.append('event_date>=?');params.append(start)
    if end: clauses.append('event_date<=?');params.append(end)
    query='SELECT * FROM hub_events'+(' WHERE '+' AND '.join(clauses) if clauses else '')+' ORDER BY event_date,id'
    with db() as c: return jsonify([dict(row) for row in c.execute(query,params)])

@app.post('/api/events')
def add_event():
    values=event_payload(request.get_json(silent=True))
    if not values: return jsonify(error='Informe data, título e categoria válidos.'),400
    with db() as c: eid=c.execute('INSERT INTO hub_events(event_date,title,category,description) VALUES(?,?,?,?)',values).lastrowid
    return jsonify(id=eid),201

@app.put('/api/events/<int:eid>')
def edit_event(eid):
    data=request.get_json(silent=True); values=event_payload(data)
    if not values: return jsonify(error='Informe data, título e categoria válidos.'),400
    with db() as c:
        row=c.execute('SELECT revision,source FROM hub_events WHERE id=?',(eid,)).fetchone()
        if not row: return jsonify(error='Data não encontrada.'),404
        if row['source']!='manual': return jsonify(error='Datas importadas são atualizadas pela planilha. Crie uma data manual para ajustes.'),409
        if data.get('revision')!=row['revision']: return jsonify(error='Esta data foi alterada. Atualize a página.'),409
        if not c.execute('UPDATE hub_events SET event_date=?,title=?,category=?,description=?,revision=revision+1 WHERE id=? AND revision=?',(*values,eid,data['revision'])).rowcount: return jsonify(error='Esta data foi alterada. Atualize a página.'),409
    return jsonify(ok=True)

@app.delete('/api/events/<int:eid>')
def delete_event(eid):
    with db() as c:
        row=c.execute('SELECT source FROM hub_events WHERE id=?',(eid,)).fetchone()
        if not row: return jsonify(error='Data não encontrada.'),404
        if row['source']!='manual': return jsonify(error='Datas importadas são atualizadas pela planilha.'),409
        c.execute('DELETE FROM hub_events WHERE id=?',(eid,))
    return jsonify(ok=True)

def parse_calendar_workbook(content):
    workbook=load_workbook(io.BytesIO(content),read_only=False,data_only=True)
    if len(workbook.worksheets)>20: raise ValueError('Muitas abas na planilha.')
    events={}
    for sheet in workbook:
        if sheet.max_row>1000 or sheet.max_column>100: raise ValueError('Planilha maior que o formato esperado.')
        if str(sheet['B3'].value or '').strip().upper()!='AUGUST' or str(sheet['I3'].value or '').strip().upper()!='SEPTEMBER' or str(sheet['P3'].value or '').strip().upper()!='OCTOBER': continue
        for month,columns,rows in ((8,range(2,8),range(5,11)),(9,range(9,15),range(5,11)),(10,range(16,22),range(5,11)),(11,range(2,8),range(14,20)),(12,range(9,15),range(14,20))):
            for row in rows:
                for col in columns:
                    cell=sheet.cell(row,col); color=cell.fill.fgColor
                    if color.type!='rgb': continue
                    found=SHEET_COLORS.get(color.rgb.upper() if isinstance(color.rgb,str) else '')
                    if not found: continue
                    value=str(int(cell.value)) if isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool) and cell.value==int(cell.value) else str(cell.value or '').strip()
                    match=re.match(r'^\s*(\d{1,2})(?:\b|$)',value)
                    if not match: continue
                    try: event_date=date(2026,month,int(match.group(1))).isoformat()
                    except ValueError: continue
                    category,title=found;key=(sheet.title,event_date,category)
                    detail=value[len(match.group(1)):].strip()
                    weekday_group='seg/qua' if re.search(r'\bSQ\b',sheet.title,re.I) else 'ter/qui' if re.search(r'\bTQ\b',sheet.title,re.I) else sheet.title
                    if key not in events: events[key]=dict(event_date=event_date,title=f'{title} · {sheet.title} ({weekday_group})',category=category,details=set(),sheet=sheet.title)
                    if detail: events[key]['details'].add(detail)
    if not events:
        workbook.close()
        return []
    for day,title in [('2026-12-15','Fim das aulas — turmas de terça e quinta'),('2026-12-16','Fim das aulas — turmas de segunda e quarta')]:
        events[('encerramento',day,'encerramento')]=dict(event_date=day,title=title,category='encerramento',details=set(),sheet='encerramento')
    workbook.close()
    return [(f'calendar-2026-2:{sheet}:{day}:{category}',day,item['title'],category,' · '.join(sorted(item['details']))[:500]) for (sheet,day,category),item in sorted(events.items())]

@app.post('/api/events/import')
def import_events():
    file=request.files.get('file')
    if not file or not file.filename.lower().endswith('.xlsx'): return jsonify(error='Selecione a planilha XLSX do calendário 2026.2.'),400
    try: rows=parse_calendar_workbook(file.read())
    except Exception:
        app.logger.exception('Falha ao ler calendário')
        return jsonify(error='Não foi possível ler essa planilha.'),400
    if not rows: return jsonify(error='Nenhuma data importante foi encontrada.'),400
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        c.execute("DELETE FROM hub_events WHERE source='calendar-2026-2'")
        c.executemany("INSERT INTO hub_events(source_key,event_date,title,category,description,source) VALUES(?,?,?,?,?,'calendar-2026-2')",rows)
    return jsonify(imported=len(rows))

def class_payload(data):
    if not isinstance(data,dict): return None
    weekday,room,start_time,course,students,teacher=[data.get(k) for k in ('weekday','room','start_time','course','students','teacher')]
    if type(weekday) is not int or weekday not in (0,1,2,3,4,5) or type(room) is not int or not 1<=room<=9 or not isinstance(start_time,str) or not re.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d',start_time) or not isinstance(course,str) or not 1<=len(course.strip())<=80 or students is not None and (type(students) is not int or not 0<=students<=100) or not isinstance(teacher,str) or len(teacher)>80: return None
    return weekday,room,start_time,course.strip(),students,teacher.strip()

@app.get('/api/classes')
def list_classes():
    with db() as c: return jsonify([dict(row) for row in c.execute('SELECT * FROM hub_classes ORDER BY weekday,start_time,room')])

@app.post('/api/classes')
def add_class():
    values=class_payload(request.get_json(silent=True))
    if not values: return jsonify(error='Informe dia, sala, horário e turma válidos.'),400
    try:
        with db() as c: cid=c.execute('INSERT INTO hub_classes(weekday,room,start_time,course,students,teacher) VALUES(?,?,?,?,?,?)',values).lastrowid
    except sqlite3.IntegrityError: return jsonify(error='Já existe uma turma nesta sala e horário.'),409
    return jsonify(id=cid),201

@app.put('/api/classes/<int:cid>')
def edit_class(cid):
    data=request.get_json(silent=True);values=class_payload(data)
    if not values: return jsonify(error='Informe dia, sala, horário e turma válidos.'),400
    try:
        with db() as c:
            row=c.execute('SELECT revision FROM hub_classes WHERE id=?',(cid,)).fetchone()
            if not row: return jsonify(error='Turma não encontrada.'),404
            if data.get('revision')!=row['revision']: return jsonify(error='Esta turma foi alterada. Atualize a página.'),409
            if not c.execute('UPDATE hub_classes SET weekday=?,room=?,start_time=?,course=?,students=?,teacher=?,revision=revision+1 WHERE id=? AND revision=?',(*values,cid,data['revision'])).rowcount: return jsonify(error='Esta turma foi alterada. Atualize a página.'),409
    except sqlite3.IntegrityError: return jsonify(error='Já existe uma turma nesta sala e horário.'),409
    return jsonify(ok=True)

@app.delete('/api/classes/<int:cid>')
def delete_class(cid):
    with db() as c:
        if not c.execute('DELETE FROM hub_classes WHERE id=?',(cid,)).rowcount: return jsonify(error='Turma não encontrada.'),404
    return jsonify(ok=True)

@app.get('/api/class-map/image')
def class_map_image():
    for ext,mime in (('png','image/png'),('jpg','image/jpeg'),('webp','image/webp')):
        path=DATA/('class-map.'+ext)
        if path.exists(): return send_file(path,mimetype=mime,conditional=True)
    abort(404)

@app.post('/api/class-map/image')
def upload_class_map_image():
    file=request.files.get('file')
    if not file: return jsonify(error='Selecione uma imagem PNG, JPG ou WebP.'),400
    content=file.read()
    if content.startswith(b'\x89PNG\r\n\x1a\n'): ext='png'
    elif content.startswith(b'\xff\xd8\xff'): ext='jpg'
    elif content.startswith(b'RIFF') and content[8:12]==b'WEBP': ext='webp'
    else: return jsonify(error='Imagem inválida. Use PNG, JPG ou WebP.'),400
    target=DATA/('class-map.'+ext);temporary=DATA/'class-map.upload'
    temporary.write_bytes(content);temporary.replace(target)
    for other in ('png','jpg','webp'):
        if other!=ext: (DATA/('class-map.'+other)).unlink(missing_ok=True)
    return jsonify(ok=True)

@app.get('/api/links')
def list_links():
    with db() as c:
        return jsonify([dict(r) for r in c.execute('SELECT * FROM hub_links ORDER BY title COLLATE NOCASE, id')])

@app.get('/avisos')
def notices_page(): return send_from_directory(app.static_folder,'avisos.html')

@app.get('/api/notices')
def list_notices():
    with db() as c:
        return jsonify([dict(r) for r in c.execute('SELECT * FROM hub_notices ORDER BY id')])

@app.route('/api/notices', methods=['POST'])
@app.route('/api/notices/<int:nid>', methods=['PUT'])
def save_notice(nid=None):
    data = request.get_json(silent=True)
    if not isinstance(data,dict): return jsonify(error='Dados inválidos.'),400
    title, body, color, seconds = [data.get(k) for k in ('title','body','color','seconds')]
    if (not isinstance(title,str) or not title.strip() or len(title)>100 or
        not isinstance(body,str) or len(body)>600 or color not in ('purple','blue','coral','yellow') or
        type(seconds) is not int or not 5<=seconds<=60):
        return jsonify(error='Informe título (até 100 caracteres), texto (até 600), uma cor da paleta e tempo de 5 a 60 segundos.'),400
    with db() as c:
        if nid is None:
            nid=c.execute('INSERT INTO hub_notices(title,body,color,seconds) VALUES(?,?,?,?)',(title.strip(),body.strip(),color,seconds)).lastrowid
        else:
            row=c.execute('SELECT revision FROM hub_notices WHERE id=?',(nid,)).fetchone()
            if not row: return jsonify(error='Aviso não encontrado. Atualize a página.'),404
            if data.get('revision')!=row['revision']: return jsonify(error='Este aviso foi alterado por outra pessoa. Atualize a página antes de editar.'),409
            changed=c.execute('UPDATE hub_notices SET title=?,body=?,color=?,seconds=?,revision=revision+1 WHERE id=? AND revision=?',(title.strip(),body.strip(),color,seconds,nid,data['revision'])).rowcount
            if not changed: return jsonify(error='Este aviso foi alterado por outra pessoa. Atualize a página.'),409
    return jsonify(id=nid),201 if request.method=='POST' else 200

@app.delete('/api/notices/<int:nid>')
def delete_notice(nid):
    with db() as c:
        if not c.execute('DELETE FROM hub_notices WHERE id=?',(nid,)).rowcount:
            return jsonify(error='Aviso não encontrado.'),404
    return jsonify(ok=True)

@app.post('/api/links')
def add_link():
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict): return jsonify(error='Dados inválidos.'),400
    fields = [payload.get(k, '') for k in ('title', 'url', 'description')]
    if not all(isinstance(v, str) for v in fields): return jsonify(error='Dados inválidos.'),400
    title, url, description = [v.strip() for v in fields]
    try:
        parsed = urlsplit(url)
        valid = parsed.scheme in ('https', 'http') and bool(parsed.hostname) and not parsed.username and not any(ch.isspace() for ch in url)
    except ValueError:
        valid = False
    if not title or len(title)>120 or len(description)>500 or len(url)>2048 or not valid:
        return jsonify(error='Informe uma legenda (até 120 caracteres) e um endereço http:// ou https:// válido.'),400
    with db() as c:
        lid = c.execute('INSERT INTO hub_links(title,url,description) VALUES(?,?,?)', (title,url,description)).lastrowid
    return jsonify(id=lid),201

@app.delete('/api/links/<int:lid>')
def delete_link(lid):
    with db() as c:
        if not c.execute('DELETE FROM hub_links WHERE id=?', (lid,)).rowcount:
            return jsonify(error='Este link já foi excluído. Atualize a página.'),404
    return jsonify(ok=True)

@app.get('/api/contacts')
def contacts():
    tokens=norm(request.args.get('q','')).split()
    with db() as c: rows=c.execute('SELECT * FROM contacts').fetchall()
    result=[]
    for row in rows:
        if not all(t in row['search'] for t in tokens): continue
        d=json.loads(row['data']); d.update(id=row['id'],revision=row['revision'])
        d['checks']=[phone(p) for p in d['phones']]
        if request.args.get('issues')=='1' and not any(p['issue'] for p in d['checks']): continue
        result.append(d)
    result.sort(key=lambda d:(not norm(d['name']).startswith(norm(request.args.get('q',''))),norm(d['name'])))
    return jsonify(result)

@app.post('/api/contacts')
def create():
    d=prepare(request.get_json()); d['manual']=True; d['sources']=['Cadastro manual']
    if not d['name']: return jsonify(error='Informe nome ou telefone.'),400
    with db() as c: cid=save(c,d)
    return jsonify(id=cid)

@app.put('/api/contacts/<int:cid>')
def edit(cid):
    payload=request.get_json(); d=prepare(payload)
    if not d['name']: return jsonify(error='Informe nome ou telefone.'),400
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        row=c.execute('SELECT * FROM contacts WHERE id=?',(cid,)).fetchone()
        if not row: abort(404)
        if payload.get('revision') != row['revision']: return jsonify(error='Outro usuário alterou este contato. Feche e abra novamente antes de salvar.'),409
        d['sources']=json.loads(row['data'])['sources']; d['manual']=True
        save(c,d,cid)
    return jsonify(ok=True)

@app.post('/api/import')
def upload():
    file=request.files.get('file'); source=request.form.get('source','').strip()
    if not file or not source: return jsonify(error='Selecione o arquivo e informe a origem.'),400
    try:
        rows=parse_file(file.read(),file.filename)
        if request.form.get('preview')=='1':
            return jsonify(total=len(rows),issues=sum(any(phone(p)['issue'] for p in d['phones']) for d in rows),sample=rows[:5])
        return jsonify(import_contacts(rows,source,file.filename))
    except (ValueError,KeyError,csv.Error) as e: return jsonify(error=str(e)),400
    except Exception:
        app.logger.exception('Falha de importação')
        return jsonify(error='Não foi possível ler o arquivo. Verifique se é um CSV ou XLSX válido e não protegido.'),400

@app.get('/api/imports')
def history():
    with db() as c: return jsonify([dict(r) for r in c.execute('SELECT * FROM imports ORDER BY id DESC LIMIT 100')])

@app.get('/api/export')
def export():
    output=io.StringIO(); writer=csv.writer(output,delimiter=';')
    writer.writerow(['Nome','Telefones','E-mail','Empresa','Cargo','Setor','Observações'])
    with db() as c:
        for row in c.execute('SELECT data FROM contacts'):
            d=json.loads(row[0]); values=[d['name'],'; '.join(d['phones']),d['emails'],d['company'],d['title'],d['department'],d['notes']]
            writer.writerow(["'"+v if v.startswith(('=','+','-','@')) else v for v in values])
    return app.response_class('\ufeff'+output.getvalue(),mimetype='text/csv',headers={'Content-Disposition':'attachment; filename=agenda.csv'})

@app.post('/api/backup')
def backup():
    folder=ROOT/'backups'; folder.mkdir(exist_ok=True)
    target=folder/('agenda-'+datetime.now().strftime('%Y%m%d-%H%M%S-%f')+'.sqlite3')
    with db() as src, sqlite3.connect(target) as dest: src.backup(dest)
    return jsonify(message='Backup salvo na pasta backups do projeto.')

init()
if __name__=='__main__':
    from waitress import serve
    port=int(os.environ.get('AGENDA_PORT','4050'))
    print(f'Agenda disponível em http://localhost:{port} | Banco: {DATA}',flush=True)
    serve(app,host='0.0.0.0',port=port,threads=6)
