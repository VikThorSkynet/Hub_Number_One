import os
import tempfile
import unittest
import io
from concurrent.futures import ThreadPoolExecutor
os.environ['AGENDA_DATA']=tempfile.mkdtemp(prefix='agenda-test-')
import app
from openpyxl import Workbook

class AgendaTests(unittest.TestCase):
    def test_workflow(self):
        rows=app.parse_file('Nome;Telefone;Empresa\nJosé Silva;11987654321;Exemplo'.encode(),'teste.csv')
        self.assertEqual(app.import_contacts(rows,'Teste','a.csv')['added'],1)
        self.assertEqual(app.import_contacts(rows,'Teste','a.csv')['skipped'],1)
        self.assertEqual(app.import_contacts(rows,'Outra conta','b.csv')['updated'],1)
        client=app.app.test_client()
        d=client.get('/api/contacts?q=jose').json[0]
        self.assertEqual(len(d['sources']),2)
        d['name']='José editado'
        self.assertEqual(client.put('/api/contacts/'+str(d['id']),json=d).status_code,200)
        self.assertEqual(client.put('/api/contacts/'+str(d['id']),json=d).status_code,409)
        app.import_contacts(rows,'Teste','a.csv')
        self.assertEqual(client.get('/api/contacts?q=editado').json[0]['name'],'José editado')
        other=app.prepare(dict(name='Outro',phones=['1234']))
        app.import_contacts([other],'Teste','novo.csv')
        self.assertEqual(len(client.get('/api/contacts').json),2)
        self.assertEqual(len(client.get('/api/contacts?issues=1').json),1)
        self.assertEqual(client.post('/api/contacts',json={'name':'x'},headers={'Origin':'https://other.example'}).status_code,403)
        self.assertEqual(client.get('/').status_code,200)
        with ThreadPoolExecutor(max_workers=5) as pool:
            statuses=list(pool.map(lambda i:app.app.test_client().post('/api/contacts',json={'name':f'Pessoa {i}'}).status_code,range(5)))
        self.assertEqual(statuses,[200]*5)
        example=client.post('/api/contacts',json={'name':'PA joão pai pyetra'}).json['id']
        for query in ['pye','yetr','PYE','joao pye','pye pai']:
            with self.subTest(query=query):
                found=client.get('/api/contacts',query_string={'q':query}).json
                self.assertIn(example,[d['id'] for d in found])

    def test_phones(self):
        for value in ['+55 (11) 98765-4321','11987654321','987654321']:
            self.assertFalse(app.phone(value)['issue'])
        for value in ['+55 (11) 3456-7890','1134567890','34567890','1234']:
            self.assertTrue(app.phone(value)['issue'])
        self.assertEqual(app.phone('55987654321')['local'],'987654321')

    def test_xlsx(self):
        book=Workbook(); sheet=book.active
        sheet.append(['Nome','Telefone']);sheet.append(['Teste XLSX',11987654321])
        out=io.BytesIO();book.save(out)
        self.assertEqual(app.parse_file(out.getvalue(),'a.xlsx')[0]['phones'],['11987654321'])


class HubTests(unittest.TestCase):
    def test_pages_and_shared_links(self):
        client=app.app.test_client()
        for path in ['/', '/lista-telefonica', '/calculadora', '/links']:
            response=client.get(path)
            self.assertEqual(response.status_code,200)
            response.close()
        created=client.post('/api/links',json={'title':'Vendas','url':'https://docs.google.com/spreadsheets/d/exemplo','description':'Matrículas'})
        self.assertEqual(created.status_code,201)
        lid=created.json['id']
        self.assertIn(lid,[r['id'] for r in app.app.test_client().get('/api/links').json])
        for url in ['javascript:alert(1)','file:///C:/test','https://','https://bad host','https://user:pass@example.com']:
            self.assertEqual(client.post('/api/links',json={'title':'Inválido','url':url}).status_code,400)
        self.assertEqual(client.post('/api/links',json=[]).status_code,400)
        self.assertEqual(client.delete('/api/links/'+str(lid),headers={'Origin':'https://other.example'}).status_code,403)
        self.assertEqual(client.delete('/api/links/'+str(lid)).status_code,200)
        self.assertEqual(client.delete('/api/links/'+str(lid)).status_code,404)

class NoticeTests(unittest.TestCase):
    def test_notice_lifecycle(self):
        client=app.app.test_client()
        original={'title':'Halloween chegando','body':'Prepare sua fantasia!\nDia 31.','color':'coral','seconds':5}
        response=client.post('/api/notices',json=original)
        self.assertEqual(response.status_code,201)
        nid=response.json['id']
        saved=next(n for n in app.app.test_client().get('/api/notices').json if n['id']==nid)
        self.assertEqual(saved['body'],original['body'])
        saved.update(title='Provas chegando',color='yellow',seconds=12)
        self.assertEqual(client.put('/api/notices/'+str(nid),json=saved).status_code,200)
        self.assertEqual(client.put('/api/notices/'+str(nid),json=saved).status_code,409)
        for patch in [{'title':' '},{'title':'a'*101},{'body':'a'*601},{'color':'red'},{'seconds':4},{'seconds':61},{'seconds':True},{'seconds':8.5}]:
            self.assertEqual(client.post('/api/notices',json={**original,**patch}).status_code,400)
        self.assertEqual(client.post('/api/notices',json=[]).status_code,400)
        for method in ['post','put','delete']:
            route='/api/notices'+('' if method=='post' else '/'+str(nid))
            self.assertEqual(getattr(client,method)(route,json=original,headers={'Origin':'https://other.example'}).status_code,403)
        self.assertEqual(client.delete('/api/notices/'+str(nid)).status_code,200)
        self.assertEqual(client.delete('/api/notices/'+str(nid)).status_code,404)
        with client.get('/avisos') as response: self.assertEqual(response.status_code,200)
        with client.get('/') as response:
            html=response.get_data(as_text=True)
            self.assertLess(html.index('hub-calendar'),html.index('Ferramentas da equipe'))

class CalendarAndClassTests(unittest.TestCase):
    def test_event_period(self):
        client=app.app.test_client()
        item={'event_date':'2027-01-30','end_date':'2027-02-03','title':'Semana de provas','category':'prova','description':''}
        self.assertEqual(client.post('/api/events',json={**item,'end_date':'2027-01-29'}).status_code,400)
        self.assertEqual(client.post('/api/events',json={**item,'end_date':'2027-02-30'}).status_code,400)
        response=client.post('/api/events',json=item)
        self.assertEqual(response.status_code,201)
        eid=response.json['id']
        rows=client.get('/api/events?from=2027-02-01&to=2027-02-02').json
        saved=next(row for row in rows if row['id']==eid)
        self.assertEqual(saved['end_date'],'2027-02-03')
        self.assertFalse(any(row['id']==eid for row in client.get('/api/events?from=2027-02-04').json))
        self.assertEqual(client.put('/api/events/'+str(eid),json={**saved,'end_date':'2027-02-05'}).status_code,200)
        self.assertEqual(client.put('/api/events/'+str(eid),json=saved).status_code,409)
        self.assertEqual(client.delete('/api/events/'+str(eid)).status_code,200)

    def test_calendar_distinguishes_weekday_groups_and_preserves_manual_event(self):
        from openpyxl.styles import PatternFill
        book=Workbook();first=book.active;first.title='CH1 - SQ 1715'
        second=book.create_sheet('GOB 2 - TQ 0930')
        for sheet in [first,second]:
            sheet['B3']='AUGUST';sheet['I3']='SEPTEMBER';sheet['P3']='OCTOBER'
            sheet['K5']='2 2A';sheet['K5'].fill=PatternFill(fill_type='solid',fgColor='FFCCFFFF')
        buf=io.BytesIO();book.save(buf)
        parsed=app.parse_calendar_workbook(buf.getvalue())
        same_day=[r for r in parsed if r[1]=='2026-09-02']
        self.assertEqual(len(same_day),2)
        self.assertTrue(any('seg/qua' in row[2] for row in same_day))
        self.assertTrue(any('ter/qui' in row[2] for row in same_day))
        self.assertEqual(app.parse_calendar_workbook(self.empty_book()),[])
        client=app.app.test_client()
        manual=client.post('/api/events',json={'event_date':'2026-09-23','title':'Reunião','category':'outro','description':''})
        self.assertEqual(manual.status_code,201)
        eid=manual.json['id']
        response=client.post('/api/events/import',data={'file':(io.BytesIO(buf.getvalue()),'calendario.xlsx')})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json['imported'],4)
        self.assertEqual(len(client.get('/api/events?from=2026-09-23&to=2026-09-29').json),1)
        self.assertEqual(client.get('/api/events?from=2026-09-02&to=2026-09-02').status_code,200)
        self.assertEqual(len(client.get('/api/events?from=2026-09-02&to=2026-09-02').json),2)
        self.assertEqual(client.post('/api/events/import',data={'file':(io.BytesIO(buf.getvalue()),'calendario.xlsx')}).json['imported'],4)
        self.assertEqual(len(client.get('/api/events').json),5)
        saved=next(e for e in client.get('/api/events').json if e['id']==eid)
        self.assertEqual(client.put('/api/events/'+str(eid),json={**saved,'title':'Nova reunião'}).status_code,200)
        self.assertEqual(client.put('/api/events/'+str(eid),json={**saved,'title':'Conflito'}).status_code,409)
        self.assertEqual(client.delete('/api/events/'+str(eid)).status_code,200)
        self.assertEqual(client.get('/api/events?from=bad').status_code,400)

    @staticmethod
    def empty_book():
        book=Workbook();buf=io.BytesIO();book.save(buf);return buf.getvalue()

    def test_class_schedule_and_map_image(self):
        client=app.app.test_client()
        item={'weekday':1,'room':2,'start_time':'17:15','course':'GOB2','students':9,'teacher':''}
        created=client.post('/api/classes',json=item)
        self.assertEqual(created.status_code,201)
        cid=created.json['id']
        self.assertEqual(client.post('/api/classes',json=item).status_code,409)
        saved=next(c for c in client.get('/api/classes').json if c['id']==cid)
        self.assertEqual(client.put('/api/classes/'+str(cid),json={**saved,'teacher':'Professora A'}).status_code,200)
        self.assertEqual(client.put('/api/classes/'+str(cid),json=saved).status_code,409)
        self.assertEqual(client.post('/api/classes',json={**item,'room':10}).status_code,400)
        self.assertEqual(client.post('/api/class-map/image',data={'file':(io.BytesIO(b'bad'),'mapa.png')}).status_code,400)
        self.assertEqual(client.post('/api/class-map/image',data={'file':(io.BytesIO(b'\x89PNG\r\n\x1a\nsmall'),'mapa.png')}).status_code,200)
        self.assertEqual(client.get('/api/class-map/image').status_code,200)
        self.assertEqual(client.delete('/api/classes/'+str(cid)).status_code,200)
        self.assertEqual(client.delete('/api/classes/'+str(cid)).status_code,404)

if __name__=='__main__':unittest.main()
